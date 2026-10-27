"""Unit tests for deterministic structured-document conversion (Part B B2.6)."""

from __future__ import annotations

import io

import openpyxl
import pytest

from kb_agent.documents.convert import (
    ConversionError,
    UnsupportedFormatError,
    convert,
)

CAP = 10 * 1024 * 1024
DEPTH = 5


def _md(fmt: str, data: bytes, cap: int = CAP) -> str:
    r = convert("f", fmt, data, max_part_bytes=cap, max_depth=DEPTH)
    assert r.part_count == 1
    return r.parts[0].markdown


# --- JSON --------------------------------------------------------------------


def test_json_array_of_objects_one_block_per_record() -> None:
    md = _md("json", b'[{"id":1},{"id":2}]')
    assert "## Record 1" in md and "## Record 2" in md
    assert "- **id**: 1" in md and "- **id**: 2" in md


def test_json_object_sections_and_nested_dotted_keys() -> None:
    md = _md("json", b'{"a":{"b":{"c":1}},"name":"x"}')
    assert "- **a.b.c**: 1" in md
    assert "- **name**: x" in md


def test_json_nested_beyond_depth_is_compact() -> None:
    # depth cap 2 -> the third level renders as compact JSON on one line
    r = convert("f", "json", b'{"a":{"b":{"c":1}}}', max_part_bytes=CAP, max_depth=2)
    md = r.parts[0].markdown
    assert '{"c":1}' in md  # compact, sorted


def test_json_scalar_arrays_joined_and_object_arrays_numbered() -> None:
    md = _md("json", b'{"tags":["x","y"],"items":[{"k":1},{"k":2}]}')
    assert "- **tags**: x; y" in md
    assert "1. k: 1" in md and "2. k: 2" in md


def test_json_literals_true_false_null() -> None:
    md = _md("json", b'{"a":true,"b":false,"c":null}')
    assert "- **a**: true" in md
    assert "- **b**: false" in md
    assert "- **c**: null" in md


def test_json_invalid_reports_location_not_content() -> None:
    with pytest.raises(ConversionError) as ei:
        convert("secret.json", "json", b'{"ssn": 123, }', max_part_bytes=CAP, max_depth=DEPTH)
    msg = str(ei.value)
    assert "line" in msg and "column" in msg
    assert "123" not in msg  # content never leaks


def test_json_bom_tolerated() -> None:
    md = _md("json", b"\xef\xbb\xbf" + b'{"a":1}')
    assert "- **a**: 1" in md


def test_jsonl_one_record_per_line_bad_line_located() -> None:
    md = _md("jsonl", b'{"id":1}\n{"id":2}\n')
    assert "## Record 1" in md and "## Record 2" in md
    with pytest.raises(ConversionError) as ei:
        convert("f", "jsonl", b'{"id":1}\n{bad}\n', max_part_bytes=CAP, max_depth=DEPTH)
    assert "line 2" in str(ei.value)


def test_json_output_determinism() -> None:
    data = b'[{"z":1,"a":{"n":2}},{"id":"a|b"}]'
    a = _md("json", data)
    b = _md("json", data)
    assert a == b


def test_json_record_boundary_split() -> None:
    data = b'[{"id":1},{"id":2},{"id":3}]'
    r = convert("big.json", "json", data, max_part_bytes=40, max_depth=DEPTH)
    assert r.part_count >= 2
    assert [p.suffix for p in r.parts] == [f"part-{i+1:03d}" for i in range(r.part_count)]


# --- CSV ---------------------------------------------------------------------


def test_csv_becomes_markdown_table() -> None:
    md = _md("csv", b"name,role\nAlice,eng\nBob,ops\n")
    assert "| name | role |" in md
    assert "| --- | --- |" in md
    assert "| Alice | eng |" in md


def test_csv_pipe_in_cell_escaped() -> None:
    md = _md("csv", b"a,b\nx|y,z\n")
    assert "x\\|y" in md


# --- XLSX --------------------------------------------------------------------


def _xlsx(sheets: dict[str, list[list]], merges: dict[str, list[str]] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=val)
        for rng in (merges or {}).get(title, []):
            ws.merge_cells(rng)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_xlsx_one_section_per_sheet() -> None:
    data = _xlsx({"S1": [["h1", "h2"], ["a", "b"]], "S2": [["x"], ["y"]]})
    md = _md("xlsx", data)
    assert "## Sheet: S1" in md and "## Sheet: S2" in md
    assert "| h1 | h2 |" in md


def test_xlsx_merged_cells_forward_filled() -> None:
    data = _xlsx({"S1": [["H", None], ["a", "b"]]}, merges={"S1": ["A1:B1"]})
    md = _md("xlsx", data)
    # the merged header value appears in both columns of the header row
    assert "| H | H |" in md


def test_xlsx_determinism() -> None:
    data = _xlsx({"S1": [["h"], ["v"]]})
    assert _md("xlsx", data) == _md("xlsx", data)


# --- .xls / unknown ----------------------------------------------------------


def test_xls_is_unsupported() -> None:
    with pytest.raises(UnsupportedFormatError):
        convert("legacy.xls", "xls", b"\xd0\xcf\x11\xe0", max_part_bytes=CAP, max_depth=DEPTH)


def test_non_utf8_reports_location_not_content() -> None:
    with pytest.raises(ConversionError) as ei:
        convert("f.csv", "csv", b"\xff\xfe\x00bad", max_part_bytes=CAP, max_depth=DEPTH)
    assert "UTF-8" in str(ei.value)
