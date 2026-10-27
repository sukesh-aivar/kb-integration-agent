"""Unit tests for detect_format (Part B B2.2): libmagic-first + extension tie-break."""

from __future__ import annotations

import io

import openpyxl

from kb_agent.documents.features import detect_format
from tests.pdf_fixtures import make_text_pdf


def test_pdf_by_magic() -> None:
    assert detect_format("a.pdf", make_text_pdf()[:2048]) == "pdf"
    # even with a misleading name, the magic bytes win for pdf
    assert detect_format("a.txt", make_text_pdf()[:2048]) == "pdf"


def test_text_family_uses_extension_tiebreak() -> None:
    body = b"name,role\nAlice,eng\n"
    assert detect_format("data.csv", body) == "csv"
    assert detect_format("notes.txt", b"just text") == "txt"
    assert detect_format("page.md", b"# Title\n") == "md"
    assert detect_format("cfg.json", b'{"a":1}') == "json"
    assert detect_format("log.jsonl", b'{"a":1}\n{"a":2}\n') == "jsonl"


def test_office_zip_disambiguated_by_extension() -> None:
    # Build a real xlsx and a real docx; magic gives the specific MIME.
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    b = io.BytesIO()
    wb.save(b)
    assert detect_format("a.xlsx", b.getvalue()) == "xlsx"

    import docx

    d = docx.Document()
    d.add_paragraph("body")
    db = io.BytesIO()
    d.save(db)
    assert detect_format("a.docx", db.getvalue()) == "docx"


def test_renamed_legacy_office_extension_is_authoritative() -> None:
    # A text-looking body named .xls must still be recognised as xls so the
    # router can REJECT it (rule 19), not treat it as txt.
    assert detect_format("report.xls", b"not really excel") == "xls"
    assert detect_format("old.doc", b"not really word") == "doc"


def test_images_by_magic() -> None:
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
    jpeg = b"\xff\xd8\xff\xe0" + b"\x00" * 64
    assert detect_format("a.png", png) == "png"
    assert detect_format("a.jpg", jpeg) == "jpeg"


def test_unknown_binary_without_extension_is_unknown() -> None:
    # Opaque binary with no helpful extension -> unknown (router REJECTs, rule 99).
    assert detect_format("blob", b"\x00\x01\x02\xff\xfe\xfd\x00\x99") == "unknown"


def test_unknown_extension_text_body_is_txt() -> None:
    # A readable text body with an unrecognised extension is treated as text.
    assert detect_format("file.xyz", b"plain readable text") == "txt"
