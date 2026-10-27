"""Unit tests for feature extraction + routing (spec 4.2, 7.5 one fixture per rule).

Builds real small documents in memory (a tiny pure-Python PDF writer for PDFs,
python-docx for DOCX) so feature extraction is exercised end-to-end via pdfplumber,
then routes them through the YAML rule
table and asserts the expected route + rule_id. Also drives every
routing_fixtures.yaml case through the router using synthetic features.
"""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml

from kb_agent.config import load_config, load_routing_rules
from kb_agent.documents.features import (
    DocumentFeatures,
    detect_mime,
    extract_features,
    normalise_mime,
)
from kb_agent.documents.router import route_document
from kb_agent.models import Route
from tests.pdf_fixtures import make_empty_pdf, make_text_pdf

REPO_ROOT = Path(__file__).resolve().parents[2]
CONFIG = load_config()
RULES = load_routing_rules()


# --- helpers to build real documents ----------------------------------------


def _make_text_pdf(text: str = "Hello world. " * 50, pages: int = 1) -> bytes:
    return make_text_pdf(text, pages)


def _make_empty_pdf(pages: int = 1) -> bytes:
    return make_empty_pdf(pages)


def _make_docx(with_image: bool = False, merged_table: bool = False) -> bytes:
    import io

    import docx

    d = docx.Document()
    d.add_paragraph("A simple paragraph of body text for the document.")
    if merged_table:
        table = d.add_table(rows=2, cols=2)
        a = table.cell(0, 0)
        b = table.cell(0, 1)
        a.merge(b)  # merged cell -> "complex"
    if with_image:
        import struct
        import zlib

        # Minimal 1x1 PNG.
        def _png() -> bytes:
            sig = b"\x89PNG\r\n\x1a\n"
            ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
            ihdr_c = b"IHDR" + ihdr
            ihdr_crc = struct.pack(">I", zlib.crc32(ihdr_c))
            ihdr_chunk = struct.pack(">I", len(ihdr)) + ihdr_c + ihdr_crc
            raw = b"\x00\xff\x00\x00"
            comp = zlib.compress(raw)
            idat_c = b"IDAT" + comp
            idat_crc = struct.pack(">I", zlib.crc32(idat_c))
            idat_chunk = struct.pack(">I", len(comp)) + idat_c + idat_crc
            iend_c = b"IEND"
            iend_chunk = struct.pack(">I", 0) + iend_c + struct.pack(">I", zlib.crc32(iend_c))
            return sig + ihdr_chunk + idat_chunk + iend_chunk

        d.add_picture(io.BytesIO(_png()))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --- MIME detection ----------------------------------------------------------


def test_detect_mime_from_bytes_not_extension() -> None:
    assert detect_mime(b"%PDF-1.4\n...") == "application/pdf"
    assert normalise_mime("application/pdf") == "pdf"


# --- real-document feature extraction + routing ------------------------------


def test_text_pdf_routes_default() -> None:
    data = _make_text_pdf()
    features = extract_features("clean.pdf", data)
    assert normalise_mime(features.detected_mime) == "pdf"
    assert features.min_page_chars and features.min_page_chars > 100
    decision = route_document(features, CONFIG, RULES)
    assert decision.route == Route.DEFAULT
    assert decision.rule_id == 10


def test_empty_pdf_routes_bda_low_text() -> None:
    data = _make_empty_pdf()
    features = extract_features("scanned.pdf", data)
    decision = route_document(features, CONFIG, RULES)
    # No text layer -> below min_chars_per_page -> BDA (rule 6).
    assert decision.route == Route.BDA
    assert decision.rule_id == 6


def test_plain_text_routes_default() -> None:
    features = extract_features("notes.txt", b"just some plain text content here")
    decision = route_document(features, CONFIG, RULES)
    assert decision.route == Route.DEFAULT
    assert decision.rule_id == 1


def test_docx_with_image_routes_bda() -> None:
    data = _make_docx(with_image=True)
    features = extract_features("withimg.docx", data)
    assert normalise_mime(features.detected_mime) == "docx"
    assert features.embedded_images > 0
    decision = route_document(features, CONFIG, RULES)
    assert decision.route == Route.BDA
    assert decision.rule_id == 3


def test_plain_docx_routes_default() -> None:
    data = _make_docx(with_image=False, merged_table=False)
    features = extract_features("plain.docx", data)
    decision = route_document(features, CONFIG, RULES)
    assert decision.route == Route.DEFAULT
    assert decision.rule_id == 4


# --- every routing fixture via synthetic features ----------------------------


def _features_from_fixture(fx: dict) -> DocumentFeatures:
    f = fx["features"]
    fmt = f["format"]
    # A representative detected_mime for the format (the router keys on `format`,
    # but the field is non-optional and surfaced in the audit dump).
    mime = {
        "txt": "text/plain", "md": "text/markdown", "html": "text/html",
        "csv": "text/csv", "json": "application/json", "jsonl": "application/json",
        "png": "image/png", "jpeg": "image/jpeg",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "doc": "application/msword",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsm": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xls": "application/vnd.ms-excel",
        "pdf": "application/pdf", "unknown": "application/octet-stream",
    }.get(fmt, "application/octet-stream")
    return DocumentFeatures(
        file_id=fx["name"],
        detected_mime=mime,
        format=fmt,
        size_bytes=f.get("size_bytes", 1000),
        pdf_encrypted_or_unreadable=f.get("pdf_encrypted_or_unreadable", False),
        min_page_chars=f.get("min_page_chars"),
        max_page_image_area_ratio=f.get("max_page_image_area_ratio", 0.0),
        detected_tables=f.get("detected_tables", 0),
        multi_column=f.get("multi_column", False),
        embedded_images=f.get("embedded_images", 0),
        complex_tables=f.get("complex_tables", 0),
    )


def _load_fixtures() -> list[dict]:
    data = yaml.safe_load((REPO_ROOT / "evals" / "routing_fixtures.yaml").read_text())
    return data["fixtures"]


@pytest.mark.parametrize("fx", _load_fixtures(), ids=lambda fx: fx["name"])
def test_routing_fixture(fx: dict) -> None:
    features = _features_from_fixture(fx)
    decision = route_document(features, CONFIG, RULES)
    assert decision.route.value == fx["expected_route"], f"{fx['name']}: route"
    assert decision.rule_id == fx["expected_rule_id"], f"{fx['name']}: rule_id"
